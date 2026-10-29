"""Local-only ingestion, transformation and reporting. No network acquisition."""
from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.runtime'))
import csv, hashlib, json, traceback, warnings
from datetime import datetime
from zoneinfo import ZoneInfo
import pandas as pd
import openpyxl

FILES = ['01_ECB_IVF_Total_Assets_2018_Onward.csv',
         '02_ECB_IVF_Fund_Shares_and_Flows_2018_Onward.csv',
         '03_CSSF_Basic_Statistical_Data_UCIs_August_2026.xlsx']
AREAS = {'DE':'Germany','FR':'France','IE':'Ireland','LU':'Luxembourg','U2':'Euro area (changing composition)'}
MEASURES = {('T00','1'):'assets',('L30','1'):'shares',('L30','2'):'gross_issues',('L30','3'):'redemptions',('L30','4'):'transactions'}

def save_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str), encoding='utf-8')

def write_table(frame, directory, name):
    directory.mkdir(parents=True, exist_ok=True)
    frame.to_parquet(directory / (name+'.parquet'), index=False)
    frame.to_csv(directory / (name+'.csv'), index=False, encoding='utf-8-sig')

def inspect_sources():
    inventory=[]; frames=[]
    for name in FILES:
        f=ROOT/'data/raw'/name
        if not f.is_file(): raise ValueError('Missing required file: '+name)
        content=f.read_bytes()
        if b'<html' in content[:1000].lower() or b'<!doctype' in content[:1000].lower():
            raise ValueError('HTML instead of source data: '+name)
        entry={'file':name,'bytes':len(content),'sha256':hashlib.sha256(content).hexdigest()}
        if f.suffix=='.csv':
            d=pd.read_csv(f,dtype=str,keep_default_na=False)
            required={'KEY','TIME_PERIOD','OBS_VALUE','REF_AREA','FREQ','IVF_ITEM','DATA_TYPE','UNIT','UNIT_MULT','TITLE_COMPL'}
            if not required <= set(d): raise ValueError('Incompatible CSV schema: '+name)
            if d.duplicated(['KEY','TIME_PERIOD']).any(): raise ValueError('Duplicate observation grain: '+name)
            if not d.TIME_PERIOD.str.fullmatch(r'\d{4}-Q[1-4]').all(): raise ValueError('Invalid quarterly periods')
            if set(d.REF_AREA)!=set(AREAS) or set(d.FREQ)!={'Q'} or set(d.UNIT)!={'EUR'} or set(d.UNIT_MULT)!={'6'}:
                raise ValueError('Unexpected geography/frequency/unit contract')
            for col,val in {'ADJUSTMENT':'N','IVF_REP_SECTOR':'T0','MATURITY_ORIG':'A','COUNT_AREA':'Z5','BS_COUNT_SECTOR':'0000','CURRENCY_TRANS':'Z01','BS_SUFFIX':'E'}.items():
                if set(d[col])!={val}: raise ValueError('Changed series definition: '+col)
            expected={('T00','1')} if name==FILES[0] else {('L30',str(i)) for i in range(1,5)}
            if set(zip(d.IVF_ITEM,d.DATA_TYPE))!=expected: raise ValueError('Unexpected statistical measures')
            pd.to_numeric(d.OBS_VALUE.replace('',None), errors='raise')
            entry.update(rows=len(d),columns=list(d),blank_counts=(d=='').sum().to_dict(),
                         period_start=d.TIME_PERIOD.min(),period_end=d.TIME_PERIOD.max(),
                         areas=sorted(d.REF_AREA.unique()),frequency=sorted(d.FREQ.unique()),
                         duplicate_keys=int(d.duplicated(['KEY','TIME_PERIOD']).sum()),
                         statuses=d.OBS_STATUS.value_counts(dropna=False).to_dict(),
                         series=d[['KEY','TITLE','TITLE_COMPL','UNIT','UNIT_MULT','COLLECTION','OBS_CONF']].drop_duplicates().to_dict('records'))
            entry['gaps']={k: sorted(set(pd.period_range(g.TIME_PERIOD.min(),g.TIME_PERIOD.max(),freq='Q').astype(str).str.replace('Q','-Q'))-set(g.TIME_PERIOD)) for k,g in d.groupby('KEY')}
            d['source_file']=name; d['source_row']=range(2,len(d)+2); frames.append(d)
        else:
            with warnings.catch_warnings(record=True) as caught:
                w=openpyxl.load_workbook(f,read_only=False,data_only=True)
            entry['reader_warnings']=[str(x.message) for x in caught]
            entry['worksheets']=[{'sheet':s.title,'max_row':s.max_row,'max_column':s.max_column,
                  'nonempty_cells':sum(v is not None for row in s.values for v in row)} for s in w]
            if len(w.worksheets)!=2: raise ValueError('CSSF worksheet count changed')
            if w.worksheets[0]['A1'].value!='STATISTIQUES OPC': raise ValueError('CSSF title changed')
            entry['report_date']=w.worksheets[0]['A2'].value
        inventory.append(entry)
    return inventory,frames,w

def transform(frames,workbook,run):
    stage=run/'staging'; out=run/'analytical'
    raw=pd.concat(frames,ignore_index=True); write_table(raw,stage,'ecb_source')
    fact=raw.copy()
    fact['reporting_area']=fact.REF_AREA.map(AREAS)
    fact['measure']=[MEASURES[(a,b)] for a,b in zip(fact.IVF_ITEM,fact.DATA_TYPE)]
    fact['measure_type']=fact.DATA_TYPE.map({'1':'stock','2':'flow','3':'flow','4':'flow'})
    fact['observation_value']=pd.to_numeric(fact.OBS_VALUE.replace('',None),errors='raise')
    fact['period_end']=pd.PeriodIndex(fact.TIME_PERIOD,freq='Q').end_time.normalize()
    fact['source_authority']='ECB'; fact['population']='ECB investment funds excluding MMFs and pension funds'
    write_table(fact,out,'fact_investment_fund_statistics')
    periods=sorted(fact.TIME_PERIOD.unique())
    write_table(pd.DataFrame({'period':periods,'period_end':pd.PeriodIndex(periods,freq='Q').end_time.normalize()}),out,'dim_period')
    write_table(pd.DataFrame([{'area_code':k,'area':v,'is_aggregate':k=='U2'} for k,v in AREAS.items()]),out,'dim_geography')
    write_table(fact[['IVF_ITEM','DATA_TYPE','measure','measure_type','TITLE','UNIT','UNIT_MULT']].drop_duplicates(),out,'dim_statistical_measure')
    # Pivot retains unavailable cells; comparisons use exact period keys, never positional lags.
    kpi=fact.pivot(index=['REF_AREA','TIME_PERIOD'],columns='measure',values='observation_value').reset_index()
    lookup=kpi.set_index(['REF_AREA','TIME_PERIOD'])
    prior=lambda a,p,n,m: lookup[m].get((a,str(pd.Period(p,freq='Q')-n).replace('Q','-Q')),float('nan'))
    for m,n,col in [('assets',1,'previous_assets'),('assets',4,'year_ago_assets'),('shares',1,'previous_shares')]:
        kpi[col]=[prior(a,p,n,m) for a,p in zip(kpi.REF_AREA,kpi.TIME_PERIOD)]
    for col,base in [('qoq_growth_pct','previous_assets'),('yoy_growth_pct','year_ago_assets')]:
        kpi[col]=(kpi.assets/kpi[base].where(kpi[base]>0)-1)*100
    kpi['net_issuance']=kpi.gross_issues-kpi.redemptions
    kpi['issuance_transaction_difference']=kpi.net_issuance-kpi.transactions
    kpi['net_issuance_rate_pct']=kpi.net_issuance/kpi.previous_shares.where(kpi.previous_shares>0)*100
    kpi['asset_change']=kpi.assets-kpi.previous_assets
    kpi['shares_change']=[v-prior(a,p,1,'shares') for a,p,v in zip(kpi.REF_AREA,kpi.TIME_PERIOD,kpi.shares)]
    kpi['shares_nontransaction_residual']=kpi.shares_change-kpi.transactions
    # This residual is not a pure valuation or investment-return measure.
    ea=kpi[kpi.REF_AREA=='U2'].set_index('TIME_PERIOD').assets
    kpi['share_of_euro_area_pct']=kpi.assets/kpi.TIME_PERIOD.map(ea).where(kpi.TIME_PERIOD.map(ea)>0)*100
    previous_net=[prior(a,p,1,'gross_issues')-prior(a,p,1,'redemptions') for a,p in zip(kpi.REF_AREA,kpi.TIME_PERIOD)]
    kpi['flow_reversal']=[None if pd.isna(v) or pd.isna(n) else bool(v*n<0) for v,n in zip(previous_net,kpi.net_issuance)]
    kpi['asset_decline']=kpi.qoq_growth_pct.map(lambda v: None if pd.isna(v) else bool(v<0))
    kpi['two_quarter_asset_decline']=[None if pd.isna(prior(a,p,2,'assets')) else bool(v<0 and prior(a,p,1,'assets')<prior(a,p,2,'assets')) for a,p,v in zip(kpi.REF_AREA,kpi.TIME_PERIOD,kpi.qoq_growth_pct)]
    kpi['area']=kpi.REF_AREA.map(AREAS); kpi['unit']='EUR millions'; kpi['population']='ECB non-MMF investment funds'
    write_table(kpi,out,'mart_ecb_quarterly')
    # Preserve every nonblank workbook cell separately, including footnotes and cached formula values.
    cells=[]
    for sheet in workbook:
        for row in sheet:
            for c in row:
                if c.value is not None: cells.append({'sheet':sheet.title,'cell':c.coordinate,'value':str(c.value),'type':type(c.value).__name__,'source_file':FILES[2]})
    write_table(pd.DataFrame(cells),stage,'cssf_source_cells')
    s=workbook.worksheets[0]; history=[]
    for row in s:
        if isinstance(row[0].value,datetime) and row[0].row>5:
            vals=[c.value for c in row[:8]]
            history.append(dict(zip(['month','uci_count','gross_issues','redemptions','net_issuance','market_variation','net_assets','monthly_change'],vals),source_row=row[0].row))
    hist=pd.DataFrame(history); hist['month']=pd.to_datetime(hist.month).dt.strftime('%Y-%m')
    hist['unit']='EUR billions except uci_count'; hist['source_file']=FILES[2]; hist['population']='CSSF supervised UCI statistics'; hist['source_sheet']=s.title
    if len(hist)==0 or hist.month.duplicated().any(): raise ValueError('Invalid CSSF monthly history')
    write_table(hist,out,'fact_cssf_monthly')
    s=workbook.worksheets[1]; structure=[]
    if s['A5'].value!="1. Nombre d'OPC" or s['A15'].value!='2. Actifs nets des OPC (en mia EUR)': raise ValueError('CSSF structural header changed')
    for start,measure,unit in [(8,'uci_count','count'),(18,'net_assets','EUR billions')]:
        for r in range(start,start+4):
            for c in range(2,6):
                structure.append({'month':s['A2'].value.strftime('%Y-%m'),'legal_status':s.cell(r,1).value,
                  'legal_form':s.cell(start-1,c).value.replace('\n',' '),'measure':measure,'value':s.cell(r,c).value,
                  'unit':unit,'source_file':FILES[2],'source_sheet':s.title,'source_cell':s.cell(r,c).coordinate,'population':'CSSF supervised UCI statistics'})
    write_table(pd.DataFrame(structure),out,'fact_cssf_structure')
    return fact,kpi,hist

def source_docs(inventory):
    docs=ROOT/'docs/data_sources'; docs.mkdir(parents=True,exist_ok=True)
    save_json(docs/'source_inventory.json',inventory)
    lines=['# Source inventory','', 'Inspected locally on 28 October 2026. Files remain unchanged; no raw data acquired. Hashes identify this supplied vintage, not remote provenance certification.','']
    schema=['# Physical source schema','']
    for e in inventory:
        lines += ['## '+e['file'],f"- Size: {e['bytes']:,} bytes",f"- SHA-256: `{e['sha256']}`"]
        if 'rows' in e:
            lines += [f"- {e['rows']} observations × {len(e['columns'])} fields; {len(e['series'])} series",f"- Coverage: {e['period_start']}–{e['period_end']}; areas {', '.join(e['areas'])}; quarterly",f"- Duplicate KEY/TIME_PERIOD: {e['duplicate_keys']}; missing OBS_VALUE: {e['blank_counts']['OBS_VALUE']}; statuses: {e['statuses']}",'- Units: EUR, UNIT_MULT=6 (EUR millions); no metadata/footer records encountered.','- No internal quarter gaps in the supplied series. Blank metadata counts and full series descriptions: source_inventory.json.','']
            schema += ['## '+e['file'],'','| Physical column | Blank records |','|---|---:|']+[f"| {c} | {e['blank_counts'][c]} |" for c in e['columns']]
        else:
            lines += ['- Genuine readable XLSX, two sheets; report date 2026-08-01 (August reference month).','- Evolution générale: A5:H5 headers, A6:H62 observations, 57 months from 2021-12 to 2026-08; A63 redemption footnote. Declared used range A1:H80 includes formatting-only trailing rows.','- Stat. détaillée: A1:H42 used range; counts A7:F12, assets A17:F22, umbrella structure A27:F28, flows A33:F36, distributions A41:F42.','- First history row has no market variation or monthly change; retain these two blanks. No duplicate months.','- Reader warnings concern print-area and header/footer formatting; all observation cells are readable.','- Redemption footnote: includes UCIs liquidated during the reference month.','']
            schema += ['## CSSF workbook','', 'History A=month, B=UCI count, C=gross issues, D=redemptions, E=net issuance, F=financial-market variation, G=net assets, H=absolute monthly change. Monetary values are EUR billions; B is a count. H is not a percentage.','', 'Detailed tables retain legal status (Part I, Part II, FIS, SICAR) and legal form (FCP, SICAV, Other UCI/FIS, SICAR). Totals are controls, excluded from the additive leaf fact. The 32-cell structural fact is 4 statuses × 4 forms × 2 measures. Umbrella counts, compartments and distributions remain available in lossless cell staging.','']
    lines += ['## Statistical boundaries','ECB: investment funds excluding MMFs and pension funds. CSSF: workbook UCI population including displayed Part I, Part II, FIS and SICAR categories; do not assert equality with the ECB perimeter. U2 denotes changing euro-area composition. No historical revision vintage is available; changes to future supplied hashes require review.']
    schema += ['## Grain and semantics','ECB grain: KEY + TIME_PERIOD. KEY retains all source dimensions. T00/1=assets stock; L30/1=shares stock; L30/2=gross issues flow; L30/3=redemptions flow; L30/4=transactions flow. All series use N/T0/A/Z5/0000/Z01/E dimensions. No supported fund-strategy breakdown exists, so no invented fund-category dimension.','', 'Source TITLE and TITLE_COMPL are preserved verbatim with all 37 physical fields. OBS_STATUS=A for all 850 observations; OBS_CONF=F. Blank OBS_PRE_BREAK, OBS_COM and other optional metadata are not missing numerical observations.']
    (docs/'source_inventory.md').write_text('\n'.join(lines),encoding='utf-8')
    (docs/'source_schema_map.md').write_text('\n'.join(schema),encoding='utf-8')

def main():
    run_id=datetime.now(ZoneInfo('Europe/Paris')).strftime('%Y%m%dT%H%M%S_%f')
    run=ROOT/'data/runs'/run_id; run.mkdir(parents=True)
    try:
        inventory,frames,w=inspect_sources(); source_docs(inventory); save_json(run/'manifest.json',inventory)
        contract=ROOT/'docs/data_sources/approved_source_contract.json'
        if contract.exists():
            old=json.loads(contract.read_text(encoding='utf-8'))
            if [(x['file'],x['sha256']) for x in old]!=[(x['file'],x['sha256']) for x in inventory]:
                raise ValueError('Source vintage changed. Review schema, definitions and revisions, then explicitly update approved_source_contract.json; prior outputs retained.')
        else: save_json(contract,inventory)
        fact,kpi,hist=transform(frames,w,run)
        from validation.validate import validate
        result=validate(ROOT,run)
        save_json(run/'validation.json',result)
        if result['failed']: raise ValueError(f"{result['failed']} independent validation checks failed")
        from reporting import executive
        executive(ROOT,run,kpi,hist,result)
        # Last-known-good pointer changes atomically, only after all reporting succeeds.
        pointer=ROOT/'data/latest_validated.json'; tmp=pointer.with_suffix('.tmp')
        save_json(tmp,{'run_id':run_id,'relative_path':f'data/runs/{run_id}','passed':result['passed'],'failed':result['failed']})
        tmp.replace(pointer)
        save_json(run/'status.json',{'status':'PASS','run_id':run_id,'checks':result})
        print(json.dumps({'run_id':run_id,'ecb_rows':len(fact),'kpi_rows':len(kpi),'cssf_months':len(hist),'passed':result['passed'],'failed':result['failed']}))
    except Exception as exc:
        save_json(run/'status.json',{'status':'FAILED','error':str(exc)})
        traceback.print_exc(); sys.exit(1)

if __name__=='__main__': main()
